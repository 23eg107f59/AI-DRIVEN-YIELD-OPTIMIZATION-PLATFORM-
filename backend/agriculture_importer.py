import hashlib
import io
import logging
import math
import re
import sqlite3
import threading
import time
import uuid
from datetime import datetime, timedelta, timezone
from urllib.error import HTTPError, URLError
from urllib.request import Request, urlopen

from openpyxl import load_workbook

from backend.agriculture_sources import SOURCES, SOURCE_BY_ID
from backend.database import connect

logger = logging.getLogger(__name__)
SYNC_LOCK = threading.Lock()
YEAR_PATTERN = re.compile(r"^\d{4}-\d{2,4}$")
RETRYABLE_HTTP_CODES = {408, 425, 429, 500, 502, 503, 504}
REQUEST_HEADERS = {"User-Agent": "OilseedAI/1.0 (public agricultural statistics importer)"}
SYNC_LOCK_TTL_SECONDS = 15 * 60

CROP_NAMES = (
    "Rapeseed & Mustard",
    "Groundnut",
    "Soybean",
    "Sunflower",
    "Castor Seed",
    "Linseed",
    "Niger seed",
    "Sesame",
    "Safflower",
    "Total Oilseeds",
    "Edible oilseeds",
)
TOTAL_CROPS = {"Total Oilseeds", "Edible oilseeds"}
MISSING_NUMBERS = {"", "-", "NA", "N.A.", "N/A"}


class ImportSourceError(RuntimeError):
    pass


def _now():
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def _source_cell_text(value):
    if value is None:
        return ""
    return str(value).strip()


def _crop_name(value):
    text = _source_cell_text(value).casefold()
    for name in CROP_NAMES:
        if name.casefold() in text:
            return name
    return None


def _region_name(value):
    text = _source_cell_text(value)
    english = re.search(r"[A-Za-z][A-Za-z ]*$", text)
    if not english:
        return None
    name = " ".join(english.group(0).split())
    if name.casefold() == "all india":
        return "India"
    return name or None


def _number(value):
    if value is None or (isinstance(value, str) and value.strip().upper() in MISSING_NUMBERS):
        return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    if not math.isfinite(number) or number < 0:
        return None
    return number


def _is_missing_number(value):
    return value is None or (
        isinstance(value, str) and value.strip().upper() in MISSING_NUMBERS
    )


def _make_record(source_id, crop, region, year, area_raw, production_raw, yield_raw):
    area_thousand_ha = _number(area_raw)
    production_thousand_tonnes = _number(production_raw)
    source_yield = _number(yield_raw)
    if not YEAR_PATTERN.fullmatch(str(year or "").strip()):
        return None
    if area_thousand_ha is None and production_thousand_tonnes is None and source_yield is None:
        return None

    area_ha = area_thousand_ha * 1000 if area_thousand_ha is not None else None
    production_tonnes = (
        production_thousand_tonnes * 1000
        if production_thousand_tonnes is not None
        else None
    )
    calculated_yield = (
        production_tonnes * 1000 / area_ha
        if area_ha is not None and area_ha > 0 and production_tonnes is not None
        else None
    )
    discrepancy = (
        source_yield is not None
        and calculated_yield is not None
        and abs(source_yield - calculated_yield) > max(1.0, source_yield * 0.05)
    )
    actual_yield = source_yield if source_yield is not None else calculated_yield
    verification_status = (
        "discrepancy"
        if discrepancy
        else "source_reported"
        if source_yield is not None
        else "calculated"
        if calculated_yield is not None
        else "yield_unavailable"
    )

    return {
        "source_id": source_id,
        "source_name": SOURCE_BY_ID[source_id]["name"],
        "source_record_id": f"{crop}|{region}|{str(year).strip()}",
        "category": "official_regional",
        "crop": crop,
        "state": region,
        "district": None,
        "season": "Annual",
        "crop_year": str(year).strip(),
        "area_ha": area_ha,
        "production_tonnes": production_tonnes,
        "yield_kg_ha": actual_yield,
        "calculated_yield_kg_ha": calculated_yield,
        "yield_discrepancy": int(discrepancy),
        "verification_status": verification_status,
        "publication_date": None,
    }


def _load_sheet(content):
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise ImportSourceError("The downloaded XLSX file could not be opened.") from exc
    if not workbook.worksheets:
        raise ImportSourceError("The downloaded workbook contains no worksheets.")
    return workbook.worksheets[0]


def parse_cropwise_workbook(content):
    sheet = _load_sheet(content)
    if sheet.max_row < 4 or _source_cell_text(sheet["A3"].value).casefold() not in {"वर्ष year", "year"}:
        raise ImportSourceError("The ICAR crop-wise workbook has an unexpected year/header layout.")

    crop_groups = []
    for column in range(2, sheet.max_column + 1):
        crop = _crop_name(sheet.cell(2, column).value)
        if not crop:
            continue
        metrics = {
            key: column + offset
            for offset, key in enumerate(("area", "production", "yield"))
        }
        header = [_source_cell_text(sheet.cell(3, metrics[key]).value).casefold() for key in metrics]
        if not all(
            marker in value
            for marker, value in zip(("area", "production", "yield"), header)
        ):
            continue
        crop_groups.append((crop, metrics))

    if not crop_groups:
        raise ImportSourceError("No crop-wise Area/Production/Yield columns were found.")

    records = []
    rejected = 0
    for row in range(4, sheet.max_row + 1):
        year = sheet.cell(row, 1).value
        for crop, columns in crop_groups:
            if crop in TOTAL_CROPS:
                continue
            values = [sheet.cell(row, columns[key]).value for key in ("area", "production", "yield")]
            if any(not _is_missing_number(value) and _number(value) is None for value in values):
                rejected += 1
                continue
            record = _make_record(
                SOURCES[0]["source_id"],
                crop,
                "India",
                year,
                *values,
            )
            if record is None:
                if any(not _is_missing_number(value) for value in values):
                    rejected += 1
                continue
            records.append(record)
    if not records:
        raise ImportSourceError("The ICAR crop-wise workbook contained no valid oilseed records.")
    return records, rejected


def parse_statewise_soybean_workbook(content):
    sheet = _load_sheet(content)
    if sheet.max_row < 4 or _source_cell_text(sheet["A3"].value).casefold() != "वर्ष":
        raise ImportSourceError("The ICAR state-wise workbook has an unexpected year/header layout.")

    groups = []
    for column in range(2, sheet.max_column + 1, 3):
        state = _region_name(sheet.cell(2, column).value)
        if not state or state.casefold() == "others":
            continue
        metric_labels = [
            _source_cell_text(sheet.cell(3, column + offset).value).casefold()
            for offset in range(3)
        ]
        if not (
            "area" in metric_labels[0]
            and "production" in metric_labels[1]
            and "yield" in metric_labels[2]
        ):
            raise ImportSourceError("The ICAR state-wise workbook's APY columns have changed.")
        groups.append((state, column))

    if not groups:
        raise ImportSourceError("No state-wise soybean APY columns were found.")

    records = []
    rejected = 0
    source_id = SOURCES[1]["source_id"]
    for row in range(4, sheet.max_row + 1):
        year = sheet.cell(row, 1).value
        for state, first_column in groups:
            values = [sheet.cell(row, first_column + offset).value for offset in range(3)]
            if any(not _is_missing_number(value) and _number(value) is None for value in values):
                rejected += 1
                continue
            record = _make_record(source_id, "Soybean", state, year, *values)
            if record is None:
                if any(not _is_missing_number(value) for value in values):
                    rejected += 1
                continue
            records.append(record)
    if not records:
        raise ImportSourceError("The ICAR state-wise workbook contained no valid soybean records.")
    return records, rejected


PARSERS = {
    "cropwise": parse_cropwise_workbook,
    "statewise_soybean": parse_statewise_soybean_workbook,
}


def _get_source_state(database_path, source_id):
    with connect(database_path) as connection:
        row = connection.execute(
            "SELECT etag, last_modified, content_hash, record_count FROM agriculture_source_state WHERE source_id = ?",
            (source_id,),
        ).fetchone()
    return dict(row) if row else {}


def _acquire_database_lock(database_path):
    now = time.time()
    token = uuid.uuid4().hex
    with connect(database_path) as connection:
        connection.execute(
            "INSERT OR IGNORE INTO agriculture_sync_locks (singleton, token, expires_at) VALUES (1, '', 0)"
        )
        stale_cutoff = (
            datetime.now(timezone.utc) - timedelta(seconds=SYNC_LOCK_TTL_SECONDS)
        ).isoformat(timespec="seconds")
        connection.execute(
            """
            UPDATE agriculture_sync_runs
            SET status = 'failed', completed_at = ?, errors = 'Previous synchronization was interrupted.'
            WHERE status IN ('queued', 'running') AND started_at < ?
            """,
            (_now(), stale_cutoff),
        )
        cursor = connection.execute(
            "UPDATE agriculture_sync_locks SET token = ?, expires_at = ? "
            "WHERE singleton = 1 AND expires_at <= ?",
            (token, now + SYNC_LOCK_TTL_SECONDS, now),
        )
    return token if cursor.rowcount else None


def _release_database_lock(database_path, token):
    if token is None:
        return
    with connect(database_path) as connection:
        connection.execute(
            "UPDATE agriculture_sync_locks SET token = '', expires_at = 0 "
            "WHERE singleton = 1 AND token = ?",
            (token,),
        )


def _download_source(source, state, timeout):
    headers = dict(REQUEST_HEADERS)
    if state.get("etag"):
        headers["If-None-Match"] = state["etag"]
    if state.get("last_modified"):
        headers["If-Modified-Since"] = state["last_modified"]
    request = Request(source["url"], headers=headers)
    for attempt in range(3):
        try:
            with urlopen(request, timeout=timeout) as response:
                return response.read(), response.headers.get("ETag"), response.headers.get("Last-Modified")
        except HTTPError as exc:
            if exc.code == 304:
                return None, state.get("etag"), state.get("last_modified")
            if exc.code not in RETRYABLE_HTTP_CODES or attempt == 2:
                raise ImportSourceError(f"Source returned HTTP {exc.code}.") from exc
            retry_after = exc.headers.get("Retry-After")
            delay = min(float(retry_after), 5.0) if retry_after and retry_after.isdigit() else 2 ** attempt
        except (TimeoutError, URLError, ConnectionError) as exc:
            if attempt == 2:
                raise ImportSourceError(f"Source download failed: {exc.reason if isinstance(exc, URLError) else exc}") from exc
            delay = 2 ** attempt
        except OSError as exc:
            if attempt == 2:
                raise ImportSourceError(f"Source download failed: {exc}") from exc
            delay = 2 ** attempt
        time.sleep(delay)
    raise ImportSourceError("Source download failed after retries.")


def _store_source_records(database_path, source, records, etag, last_modified, content_hash, rejected):
    now = _now()
    inserted = updated = skipped = 0
    with connect(database_path) as connection:
        for record in records:
            existing = connection.execute(
                "SELECT * FROM agriculture_records WHERE source_id = ? AND source_record_id = ?",
                (source["source_id"], record["source_record_id"]),
            ).fetchone()
            values = (
                record["source_name"], record["category"], record["crop"], record["state"],
                record["district"], record["season"], record["crop_year"], record["area_ha"],
                record["production_tonnes"], record["yield_kg_ha"],
                record["calculated_yield_kg_ha"], record["yield_discrepancy"],
                record["verification_status"], record["publication_date"],
            )
            if existing is None:
                connection.execute(
                    """
                    INSERT INTO agriculture_records (
                        source_id, source_record_id, source_name, category, crop, state, district,
                        season, crop_year, area_ha, production_tonnes, yield_kg_ha,
                        calculated_yield_kg_ha, yield_discrepancy, verification_status,
                        publication_date, last_synced_at
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    (source["source_id"], record["source_record_id"], *values, now),
                )
                inserted += 1
                continue

            old_values = tuple(existing[key] for key in (
                "source_name", "category", "crop", "state", "district", "season", "crop_year",
                "area_ha", "production_tonnes", "yield_kg_ha", "calculated_yield_kg_ha",
                "yield_discrepancy", "verification_status", "publication_date",
            ))
            if old_values == values:
                connection.execute(
                    "UPDATE agriculture_records SET last_synced_at = ? WHERE id = ?",
                    (now, existing["id"]),
                )
                skipped += 1
                continue
            connection.execute(
                """
                UPDATE agriculture_records
                SET source_name = ?, category = ?, crop = ?, state = ?, district = ?, season = ?,
                    crop_year = ?, area_ha = ?, production_tonnes = ?, yield_kg_ha = ?,
                    calculated_yield_kg_ha = ?, yield_discrepancy = ?, verification_status = ?,
                    publication_date = ?, last_synced_at = ?
                WHERE source_id = ? AND source_record_id = ?
                """,
                (*values, now, source["source_id"], record["source_record_id"]),
            )
            updated += 1

        connection.execute(
            """
            INSERT INTO agriculture_source_state (
                source_id, status, error, etag, last_modified, content_hash, record_count,
                last_checked_at, last_successful_sync
            ) VALUES (?, 'connected', NULL, ?, ?, ?, ?, ?, ?)
            ON CONFLICT(source_id) DO UPDATE SET
                status = excluded.status, error = NULL, etag = excluded.etag,
                last_modified = excluded.last_modified, content_hash = excluded.content_hash,
                record_count = excluded.record_count, last_checked_at = excluded.last_checked_at,
                last_successful_sync = excluded.last_successful_sync
            """,
            (source["source_id"], etag, last_modified, content_hash, len(records), now, now),
        )
    return {"inserted": inserted, "updated": updated, "skipped": skipped, "rejected": rejected}


def _mark_source(database_path, source_id, status, error=None):
    now = _now()
    with connect(database_path) as connection:
        connection.execute(
            """
            INSERT INTO agriculture_source_state (source_id, status, error, last_checked_at)
            VALUES (?, ?, ?, ?)
            ON CONFLICT(source_id) DO UPDATE SET
                status = excluded.status, error = excluded.error, last_checked_at = excluded.last_checked_at
            """,
            (source_id, status, error, now),
        )


def _sync_source(database_path, source, timeout):
    state = _get_source_state(database_path, source["source_id"])
    content, etag, last_modified = _download_source(source, state, timeout)
    if content is None:
        now = _now()
        with connect(database_path) as connection:
            count = connection.execute(
                "SELECT COUNT(*) AS total FROM agriculture_records WHERE source_id = ?",
                (source["source_id"],),
            ).fetchone()["total"]
            connection.execute(
                """
                UPDATE agriculture_source_state SET status = 'unchanged', error = NULL,
                    last_checked_at = ?, last_successful_sync = ? WHERE source_id = ?
                """,
                (now, now, source["source_id"]),
            )
            connection.execute(
                "UPDATE agriculture_records SET last_synced_at = ? WHERE source_id = ?",
                (now, source["source_id"]),
            )
        return {"retrieved": 0, "inserted": 0, "updated": 0, "skipped": count, "rejected": 0}

    content_hash = hashlib.sha256(content).hexdigest()
    if state.get("content_hash") == content_hash:
        now = _now()
        with connect(database_path) as connection:
            count = connection.execute(
                "SELECT COUNT(*) AS total FROM agriculture_records WHERE source_id = ?",
                (source["source_id"],),
            ).fetchone()["total"]
            connection.execute(
                """
                UPDATE agriculture_source_state SET status = 'unchanged', error = NULL,
                    etag = COALESCE(?, etag), last_modified = COALESCE(?, last_modified),
                    last_checked_at = ?, last_successful_sync = ? WHERE source_id = ?
                """,
                (etag, last_modified, now, now, source["source_id"]),
            )
            connection.execute(
                "UPDATE agriculture_records SET last_synced_at = ? WHERE source_id = ?",
                (now, source["source_id"]),
            )
        return {"retrieved": count, "inserted": 0, "updated": 0, "skipped": count, "rejected": 0}

    records, rejected = PARSERS[source["parser"]](content)
    counts = _store_source_records(
        database_path, source, records, etag, last_modified, content_hash, rejected
    )
    return {"retrieved": len(records) + rejected, **counts}


def run_synchronization(
    database_path,
    timeout=20,
    acquire_lock=True,
    acquire_database_lock=True,
    existing_run_id=None,
    database_lock_token=None,
):
    if acquire_lock and not SYNC_LOCK.acquire(blocking=False):
        return None
    if acquire_database_lock:
        database_lock_token = _acquire_database_lock(database_path)
        if database_lock_token is None:
            if acquire_lock:
                SYNC_LOCK.release()
            return None
    run_id = existing_run_id
    totals = {"retrieved": 0, "inserted": 0, "updated": 0, "skipped": 0, "rejected": 0}
    source_errors = []
    started_at = _now()
    try:
        with connect(database_path) as connection:
            if run_id is None:
                cursor = connection.execute(
                    "INSERT INTO agriculture_sync_runs (status, started_at) VALUES ('running', ?)",
                    (started_at,),
                )
                run_id = cursor.lastrowid
            else:
                connection.execute(
                    "UPDATE agriculture_sync_runs SET status = 'running', started_at = ? WHERE id = ?",
                    (started_at, run_id),
                )

        for source in SOURCES:
            try:
                counts = _sync_source(database_path, source, timeout)
                for key in totals:
                    totals[key] += counts[key]
            except (ImportSourceError, OSError, sqlite3.Error, ValueError) as exc:
                message = f"{source['name']}: {exc}"
                logger.warning("Agricultural source sync failed: %s", message)
                source_errors.append(message)
                _mark_source(database_path, source["source_id"], "error", str(exc))

        status = "failed" if source_errors and totals["retrieved"] == 0 else "partial" if source_errors else "success"
        completed_at = _now()
        with connect(database_path) as connection:
            connection.execute(
                """
                UPDATE agriculture_sync_runs
                SET status = ?, completed_at = ?, records_retrieved = ?, inserted = ?,
                    updated = ?, skipped = ?, rejected = ?, errors = ?
                WHERE id = ?
                """,
                (status, completed_at, totals["retrieved"], totals["inserted"],
                 totals["updated"], totals["skipped"], totals["rejected"],
                 "\n".join(source_errors) or None, run_id),
            )
        return {"id": run_id, "status": status, **totals, "errors": source_errors}
    except sqlite3.Error:
        logger.exception("Agricultural synchronization database operation failed")
        raise
    finally:
        _release_database_lock(database_path, database_lock_token)
        if acquire_lock:
            SYNC_LOCK.release()


def start_manual_synchronization(database_path, timeout=20):
    if not SYNC_LOCK.acquire(blocking=False):
        return None
    database_lock_token = _acquire_database_lock(database_path)
    if database_lock_token is None:
        SYNC_LOCK.release()
        return None
    try:
        with connect(database_path) as connection:
            cursor = connection.execute(
                "INSERT INTO agriculture_sync_runs (status, started_at) VALUES ('queued', ?)",
                (_now(),),
            )
            run_id = cursor.lastrowid
    except Exception:
        _release_database_lock(database_path, database_lock_token)
        SYNC_LOCK.release()
        raise

    def run():
        try:
            run_synchronization(
                database_path,
                timeout=timeout,
                acquire_lock=False,
                acquire_database_lock=False,
                existing_run_id=run_id,
                database_lock_token=database_lock_token,
            )
        except Exception:
            logger.exception("Manual agricultural synchronization failed")
            with connect(database_path) as connection:
                connection.execute(
                    "UPDATE agriculture_sync_runs SET status = 'failed', completed_at = ?, errors = ? WHERE id = ?",
                    (_now(), "Synchronization failed. Check the application log.", run_id),
                )
        finally:
            _release_database_lock(database_path, database_lock_token)
            SYNC_LOCK.release()

    thread = threading.Thread(target=run, name="agriculture-manual-sync", daemon=True)
    try:
        thread.start()
    except RuntimeError:
        _release_database_lock(database_path, database_lock_token)
        SYNC_LOCK.release()
        with connect(database_path) as connection:
            connection.execute(
                "UPDATE agriculture_sync_runs SET status = 'failed', completed_at = ?, errors = ? WHERE id = ?",
                (_now(), "The synchronization worker could not be started.", run_id),
            )
        raise
    return run_id
